"""
Document loaders: extract raw text from PDF, TXT, and DOCX files.

Each loader returns a LoadedDocument. The dispatcher `load_document()`
picks the correct loader based on file extension.
"""

import os
from pathlib import Path

from pypdf import PdfReader
from docx import Document as DocxDocument

from app.models.schemas import LoadedDocument
from app.config.settings import settings
from app.utils.logger import get_logger

logger = get_logger(__name__)


class DocumentLoadError(Exception):
    """Raised when a document cannot be loaded or is invalid."""
    pass


def _validate_file(file_path: str) -> None:
    path = Path(file_path)

    if not path.exists():
        raise DocumentLoadError(f"File not found: {file_path}")

    allowed_extensions = [
        ext.strip().lower()
        for ext in settings.allowed_upload_extensions.split(",")
    ]
    if path.suffix.lower() not in allowed_extensions:
        raise DocumentLoadError(
            f"Unsupported file extension '{path.suffix}'. "
            f"Allowed: {allowed_extensions}"
        )

    size_mb = path.stat().st_size / (1024 * 1024)
    if size_mb > settings.max_upload_size_mb:
        raise DocumentLoadError(
            f"File too large: {size_mb:.2f}MB "
            f"(max {settings.max_upload_size_mb}MB)"
        )


def load_pdf(file_path: str) -> LoadedDocument:
    reader = PdfReader(file_path)
    pages_text = []

    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        pages_text.append(text)

    full_text = "\n\n".join(pages_text).strip()

    if not full_text:
        raise DocumentLoadError(
            f"No extractable text found in PDF: {file_path}. "
            "It may be a scanned/image-only PDF (OCR not implemented yet)."
        )

    return LoadedDocument(
        source_filename=Path(file_path).name,
        document_type="pdf",
        full_text=full_text,
        page_count=len(reader.pages),
        char_count=len(full_text),
    )


def load_txt(file_path: str) -> LoadedDocument:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        full_text = f.read().strip()

    if not full_text:
        raise DocumentLoadError(f"TXT file is empty: {file_path}")

    return LoadedDocument(
        source_filename=Path(file_path).name,
        document_type="txt",
        full_text=full_text,
        page_count=None,
        char_count=len(full_text),
    )


def load_docx(file_path: str) -> LoadedDocument:
    doc = DocxDocument(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    full_text = "\n\n".join(paragraphs).strip()

    if not full_text:
        raise DocumentLoadError(f"No extractable text found in DOCX: {file_path}")

    return LoadedDocument(
        source_filename=Path(file_path).name,
        document_type="docx",
        full_text=full_text,
        page_count=None,
        char_count=len(full_text),
    )


_LOADER_MAP = {
    ".pdf": load_pdf,
    ".txt": load_txt,
    ".docx": load_docx,
}


def load_document(file_path: str) -> LoadedDocument:
    """
    Unified entry point: validates the file, then dispatches to the
    correct loader based on extension.
    """
    _validate_file(file_path)

    ext = Path(file_path).suffix.lower()
    loader_fn = _LOADER_MAP.get(ext)

    if loader_fn is None:
        # Shouldn't happen since _validate_file already checked,
        # but kept as a defensive guard.
        raise DocumentLoadError(f"No loader registered for extension: {ext}")

    logger.info(f"Loading document: {file_path}")
    document = loader_fn(file_path)
    logger.info(
        f"Loaded '{document.source_filename}' "
        f"({document.char_count} chars, {document.page_count or 'N/A'} pages)"
    )
    return document